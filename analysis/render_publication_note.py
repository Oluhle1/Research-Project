"""Render an internal publication-methodology review note, not a venue paper.

The report source and any numerical interpretation are supplied separately.
This renderer performs no survey analysis and does not change the Honours files.
"""
from __future__ import annotations

import argparse
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PUBLICATION = ROOT / "paper/publication"
DEFAULT_SOURCE = PUBLICATION / "methodology_results.md"
DEFAULT_OUTPUT = PUBLICATION / "Jawe_Publication_Methodology_Revision.docx"
FONT = "Times New Roman"
PAGE_WIDTH = 8.2677
CONTENT_WIDTH = PAGE_WIDTH - 1.4


def inline(paragraph, text: str, *, size: float | None = None, bold=False) -> None:
    """Support the limited emphasis needed by the author-review Markdown."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        strong = part.startswith("**") and part.endswith("**")
        code = part.startswith("`") and part.endswith("`")
        run = paragraph.add_run(part[2:-2] if strong else part[1:-1] if code else part)
        run.font.name = FONT
        run.bold = bold or strong
        if size is not None:
            run.font.size = Pt(size)


def create_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(PAGE_WIDTH), Inches(11.6929)
    section.left_margin = section.right_margin = Inches(.70)
    section.top_margin = section.bottom_margin = Inches(.70)
    section.header_distance = section.footer_distance = Inches(.27)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(11)
    normal.paragraph_format.line_spacing = 1.05
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in [("Title", 20), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 11)]:
        style = doc.styles[name]
        style.font.name, style.font.size = FONT, Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("173D57")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)
    header = section.header.paragraphs[0]
    header.text = "Publication methodology and validation — author review"
    for run in header.runs:
        run.font.name, run.font.size = FONT, Pt(8)
        run.font.color.rgb = RGBColor.from_string("526575")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    inline(footer, "Internal review note • ", size=8)
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    footer._p.append(page)
    doc.core_properties.title = "Publication Methodology and Validation: Author Review"
    doc.core_properties.author = "Oluhle Jawe"
    doc.core_properties.subject = "Technical review report; not an anonymised conference submission"
    doc.core_properties.comments = "Generated from the separately reviewed methodology-results source."
    return doc


def widths_for(headers: list[str], rows: list[list[str]]) -> list[float]:
    """Keep numerical metric columns legible and give interpretation prose room."""
    n = len(headers)
    if not 2 <= n <= 5:
        raise ValueError("The review-note renderer supports two to five table columns")
    if n == 2:
        weights = [1.8, 5.0]
    elif n == 3:
        weights = [2.5, 2.15, 2.15]
    elif n == 4:
        weights = [2.5, 1.45, 1.45, 1.4]
    else:
        weights = [2.2, 1.15, 1.15, 1.15, 1.15]
    # Text-heavy middle columns in a feedback mapping need more width than
    # short status fields. Metric tables retain their ordinary numeric layout.
    lengths = [max(len(row[j]) for row in [headers, *rows]) for j in range(n)]
    if n == 3 and lengths[1] > 90 and lengths[2] < 55:
        weights = [1.45, 3.8, 1.55]
    total = sum(weights)
    return [CONTENT_WIDTH * weight / total for weight in weights]


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    if any(len(row) != len(headers) for row in rows):
        raise ValueError("Inconsistent Markdown table columns")
    widths = widths_for(headers, rows)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for i, values in enumerate([headers, *rows]):
        cells = table.rows[0].cells if i == 0 else table.add_row().cells
        for cell, value, width in zip(cells, values, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.03
            inline(paragraph, value, size=9.5 if len(headers) > 3 else 10, bold=i == 0)
            properties = cell._tc.get_or_add_tcPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "244D68" if i == 0 else "EFF4F7" if i % 2 else "FFFFFF")
            properties.append(shading)
            if i == 0:
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
            margins = OxmlElement("w:tcMar")
            for side in ["top", "bottom", "left", "right"]:
                margin = OxmlElement("w:" + side)
                margin.set(qn("w:w"), "50" if side in ["top", "bottom"] else "65")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            properties.append(margins)
        properties = table.rows[i]._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:cantSplit"))
        if i == 0:
            properties.append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def image_path(path_text: str) -> Path:
    relative = Path(path_text)
    if relative.is_absolute() or ".." in relative.parts:
        raise ValueError("Review-note images must use paths within paper/publication")
    path = (PUBLICATION / relative).resolve()
    if not path.is_relative_to(PUBLICATION.resolve()) or not path.is_file():
        raise ValueError(f"Missing or unsupported review-note image: {path_text}")
    return path


def render(source: Path = DEFAULT_SOURCE, output: Path = DEFAULT_OUTPUT) -> Path:
    doc = create_document()
    lines = source.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line.startswith("|"):
            block = [line]
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            cells = [[value.strip() for value in row.strip("|").split("|")] for row in block]
            if len(cells) < 2 or not all(re.fullmatch(r":?-+:?", value) for value in cells[1]):
                raise ValueError("Malformed Markdown table")
            add_table(doc, cells[0], cells[2:])
            continue
        image = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if image:
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(image_path(image[2])), width=Inches(CONTENT_WIDTH))
            caption = doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(caption, image[1], size=9)
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            count = len(heading[1])
            paragraph = doc.add_paragraph(style="Title" if count == 1 else f"Heading {count - 1}")
            inline(paragraph, heading[2])
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", line)
        if bullet:
            paragraph = doc.add_paragraph(style="List Bullet")
            inline(paragraph, bullet[1])
            continue
        paragraph = doc.add_paragraph()
        inline(paragraph, line)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


def create_workflow_figure(directory: Path | None = None) -> tuple[Path, Path]:
    """Draw the recorded procedure; the run record supplies completion status."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch

    directory = PUBLICATION / "figures" if directory is None else directory
    directory.mkdir(parents=True, exist_ok=True)
    figure, axis = plt.subplots(figsize=(7.2, 8.3))
    axis.set(xlim=(0, 1), ylim=(0, 1))
    axis.axis("off")
    panels = [
        (0.02, "A. Repeated nested validation", [
            "Frozen KZN cohort\n1,013 adults; 67 events\n827 original households",
            "20 partitions × 5 outer folds\nHouseholds: primary grouping\nDerived PSUs: sensitivity grouping",
            "Within each outer training set\n5 inner grouped folds\nTraining-only imputation and scaling\nChoose C by pooled person Brier score",
            "Refit on the outer training set\nPredict the untouched outer test fold\nPair with training-prevalence reference",
            "Pool one held-out score per adult\nAUC, AP, Brier and calibration\nPartition/refit summaries describe\nsensitivity, not confidence intervals",
        ]),
        (0.52, "B. Development bootstrap", [
            "1,000 household draws\n827 households with replacement\nRetain adults and multiplicities\nKeep copies in their original group",
            "Within each bootstrap sample\n5 inner household-grouped folds\nRepeat imputation, scaling, C selection\nand final model fitting",
            "Score bootstrap and original cohort\nOriginal-cohort evaluation overlaps\ndevelopment deliberately: it is an\noptimism-correction component",
            "Correct apparent Brier and AUC\nCorrect paired Brier improvement\nMonte Carlo error describes finite\nsimulation error, not a performance CI",
            "Out-of-bag scores: diagnostics only\nRecord support, failures and tuning\nNo external validation or causal claim\nOriginal primary remains exploratory",
        ]),
    ]
    centers = [.84, .67, .50, .33, .16]
    width, height = .46, .135
    for x, title, nodes in panels:
        axis.text(x + width / 2, .975, title, ha="center", va="top", fontsize=11.2,
                  fontweight="bold", color="#173D57")
        for i, (center, text) in enumerate(zip(centers, nodes)):
            patch = FancyBboxPatch((x, center - height / 2), width, height,
                                   boxstyle="round,pad=0.008,rounding_size=0.012",
                                   linewidth=.8, edgecolor="#557487",
                                   facecolor="#E9F1F6" if i in [0, 4] else "#F8FAFC")
            axis.add_patch(patch)
            axis.text(x + width / 2, center, text, ha="center", va="center", fontsize=9.2,
                      linespacing=1.22, color="#172E3D")
            if i:
                axis.annotate("", xy=(x + width / 2, center + height / 2 + .009),
                              xytext=(x + width / 2, centers[i - 1] - height / 2 - .009),
                              arrowprops={"arrowstyle": "-|>", "lw": .9, "color": "#557487"})
    axis.text(.5, .027, "Procedure schematic; completion and convergence are reported separately.\n"
              "C: inverse penalty strength. AP: average precision. PSU: primary sampling unit.",
              ha="center", va="center", fontsize=8.3, color="#405A6B")
    figure.subplots_adjust(left=.02, right=.98, top=.99, bottom=.01)
    png, pdf = directory / "validation_workflow.png", directory / "validation_workflow.pdf"
    figure.savefig(png, dpi=220)
    figure.savefig(pdf)
    plt.close(figure)
    return png, pdf


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--diagram", action="store_true", help="Also draw the recorded validation-workflow schematic")
    args = parser.parse_args()
    if not args.source.is_file():
        raise SystemExit(f"Report source does not exist: {args.source}")
    if args.diagram:
        for path in create_workflow_figure():
            print(path)
    print(render(args.source, args.output))


if __name__ == "__main__":
    main()

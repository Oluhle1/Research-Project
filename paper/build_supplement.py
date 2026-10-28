"""Build the supplementary Word document from the recorded Markdown evidence.

This renderer preserves literal query strings; it does not run any analysis.
"""
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Jawe_OPG_Supplementary_Material.docx"
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.2677), Inches(11.6929)
section.left_margin = section.right_margin = Inches(.70)
section.top_margin = section.bottom_margin = Inches(.65)
section.header_distance = section.footer_distance = Inches(.27)
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.05
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.widow_control = True
for name, size in [("Title", 20), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 11)]:
    st = doc.styles[name]
    st.font.name = "Times New Roman"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string("153A5B")
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.space_before = Pt(9)
    st.paragraph_format.space_after = Pt(6)
header = section.header.paragraphs[0]
header.text = "JAWE • OPG NON-RECEIPT STUDY • SUPPLEMENTARY MATERIAL"
for r in header.runs:
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string("526575")
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Supplement • ")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)


def paragraph(text, style=None):
    p = doc.add_paragraph(style=style)
    # Only paired backticks mark paths/code. Wildcards and ASCII query quotes
    # remain literal, including asterisks in database search strings.
    for part in re.split(r"(`[^`]+`)", text):
        run = p.add_run(part[1:-1] if part.startswith("`") and part.endswith("`") else part)
        if part.startswith("`") and part.endswith("`"):
            run.font.size = Pt(10)
    return p


def table(headers, rows):
    if len(headers) == 11:
        paragraph("Study group and fitting", "Heading 3")
        table(headers[:6], [row[:6] for row in rows])
        paragraph("Held-out performance (same scenario IDs)", "Heading 3")
        table([headers[0]] + headers[6:], [[row[0]] + row[6:] for row in rows])
        return
    n = len(headers)
    widths = {
        2: [5.30, 1.56],
        3: [.55, 5.20, 1.11],
        4: [.85, 3.15, 1.43, 1.43],
        5: [2.10, 1.19, 1.19, 1.19, 1.19],
        6: [1.86, 1., 1., 1., 1., 1.],
        7: [.92, .95, .95, 1.03, .95, .95, 1.11],
        8: [.40, .85, .80, .70, .70, .80, .90, 1.71],
    }[n]
    if headers == ["Reporting element", "Interpretation"]:
        widths = [1.86, 5.00]
    elif headers[0] == "Coding" and n == 8:
        widths = [1.80, 1.20, .55, .55, .66, .58, .62, .90]
    elif headers[0] == "Coding" and n == 6:
        widths = [2.16, .85, .90, .95, 1., 1.]
    elif headers[0] == "ID" and n == 6:
        if headers[1] == "Check":
            widths = [.40, 3.25, .90, .80, .80, .71]
        else:
            widths = [.55, 1.10, 1.10, 1.30, 1.20, 1.61]
    elif headers[0] == "Measure" and n == 4:
        widths = [2.16, .90, 1.90, 1.90]
    elif headers[0] == "Coefficient" and n == 5:
        widths = [2.26, .90, .90, .90, 1.90]
    assert abs(sum(widths) - 6.86) < 1e-9
    t = doc.add_table(rows=1, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for i, row in enumerate([headers] + rows):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        for cell, value, width in zip(cells, row, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(value)
            run.font.name = "Times New Roman"
            run.font.size = Pt(9.5)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            pr = cell._tc.get_or_add_tcPr()
            sh = OxmlElement("w:shd")
            sh.set(qn("w:fill"), "244D68" if i == 0 else ("F0F5F8" if i % 2 else "FFFFFF"))
            pr.append(sh)
            margins = OxmlElement("w:tcMar")
            for side in ["top", "bottom", "left", "right"]:
                item = OxmlElement("w:" + side)
                item.set(qn("w:w"), "55" if side in ["top", "bottom"] else "65")
                item.set(qn("w:type"), "dxa")
                margins.append(item)
            pr.append(margins)
        props = t.rows[i]._tr.get_or_add_trPr()
        props.append(OxmlElement("w:cantSplit"))
        if i == 0:
            props.append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


lines = (ROOT / "supplement_screening.md").read_text().splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1
        continue
    if line.startswith("|"):
        block = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [v.strip() for v in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r"[:\- ]+", v) for v in cells):
                block.append(cells)
            i += 1
        table(block[0], block[1:])
        continue
    if line.startswith("# "):
        paragraph(line[2:], "Title")
    elif line.startswith("## "):
        doc.add_page_break()
        paragraph(line[3:], "Heading 1")
    elif line.startswith("### "):
        paragraph(line[4:], "Heading 2")
    else:
        paragraph(line)
    i += 1

doc.core_properties.title = "OPG Non-Receipt Study: Supplementary Material S1–S5"
doc.core_properties.author = "Oluhle Jawe"
doc.core_properties.subject = "Screening, validation, descriptive and recorded search evidence"
doc.save(OUT)
print(OUT)
